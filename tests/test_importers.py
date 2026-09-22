from __future__ import annotations

import tempfile
import unittest
import json
import os
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from unittest import mock

from openpyxl import Workbook

from wealthfolio_importer.parsers import (
    parse_revolut_current,
    parse_revolut_savings,
    parse_revolut_stocks,
    parse_sabadell,
    parse_sabadell_card,
    parse_xtb,
)
from wealthfolio_importer.mcp import McpError, McpImportResult, activity_to_mcp
from wealthfolio_importer.model import Activity
from wealthfolio_importer.service import WatchService


class ImporterTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def write(self, name: str, text: str) -> Path:
        path = self.root / name
        path.write_text(text, encoding="utf-8")
        return path

    def test_revolut_stocks_maps_cash_trade_and_tax_correction(self):
        path = self.write(
            "stocks.tsv",
            "Date\tTicker\tType\tQuantity\tPrice per share\tTotal Amount\tCurrency\tFX Rate\n"
            "2025-01-01T10:00:00Z\t\tCASH TOP-UP\t\t\tUSD 100\tUSD\t1.25\n"
            "2025-01-01T11:00:00Z\tACME\tBUY - MARKET\t2\tUSD 40\tUSD 80\tUSD\t1.25\n"
            "2025-02-01T11:00:00Z\tACME\tDIVIDEND TAX (CORRECTION)\t\t\tUSD -0.30\tUSD\t1.20\n"
            "2025-02-01T11:00:01Z\tACME\tDIVIDEND TAX (CORRECTION)\t\t\tUSD 0.30\tUSD\t1.20\n",
        )
        result = parse_revolut_stocks(path, {"currency": "USD"})
        self.assertEqual([a.activity_type for a in result.activities], ["DEPOSIT", "BUY", "TAX", "CREDIT"])
        self.assertEqual(result.activities[0].fx_rate, Decimal("0.8"))
        self.assertEqual(result.activities[-1].subtype, "REIMBURSEMENT")

    def test_revolut_current_reconciles_and_excludes_internal_transfers_from_spending(self):
        path = self.write(
            "account.tsv",
            "Tipo\tProducto\tFecha de inicio\tFecha de finalizaciÃ³n\tDescripciÃ³n\t"
            "Importe\tComisiÃ³n\tDivisa\tState\tSaldo\n"
            "Transferir\tActual\t2026-09-03 0:00:11\t2026-09-03 0:00:11\t"
            "To Francisco Javier Caro Munar\t-248.04\t0\tEUR\tCOMPLETADO\t299.13\n"
            "Transferir\tActual\t2026-09-07 9:58:16\t2026-09-07 9:58:16\t"
            "Desde EUR Ahorro Euro\t1000\t0\tEUR\tCOMPLETADO\t1299.13\n"
            "Cambio\tActual\t2026-09-07 9:58:52\t2026-09-07 9:58:52\t"
            "ConversiÃ³n a USD\t-648.22\t0\tEUR\tCOMPLETADO\t650.91\n"
            "Pago con tarjeta\tActual\t2026-09-17 9:08:18\t2026-09-17 9:08:18\t"
            "Google Cloud\t-5\t0\tEUR\tCOMPLETADO\t645.91\n"
            "Transferir\tActual\t2026-09-17 9:08:19\t2026-09-17 9:08:19\t"
            "A EUR Ahorro Euro\t-1\t0\tEUR\tCOMPLETADO\t644.91\n",
        )
        result = parse_revolut_current(
            path,
            {
                "currency": "EUR",
                "wealthfolioAccountId": "current-account-id",
            },
        )
        self.assertEqual(result.checks["openingBalance"], "547.17")
        self.assertEqual(result.checks["statementBalance"], "644.91")
        self.assertEqual(
            result.checks["activityCounts"],
            {
                "DEPOSIT": 1,
                "WITHDRAWAL": 2,
                "TRANSFER_IN": 1,
                "TRANSFER_OUT": 2,
            },
        )
        self.assertIn("Google Cloud", next(a.comment for a in result.activities if a.amount == 5))

        repeated = parse_revolut_current(path, {"currency": "EUR", "wealthfolioAccountId": "current-account-id"})
        self.assertEqual(
            [activity.identifier for activity in result.activities],
            [activity.identifier for activity in repeated.activities],
        )

    def test_revolut_current_skips_pending_rows_and_reconciles_fees(self):
        path = self.write(
            "account.tsv",
            "Tipo\tProducto\tFecha de inicio\tFecha de finalización\tDescripción\t"
            "Importe\tComisión\tDivisa\tState\tSaldo\n"
            "Transferir\tActual\t2026-09-01 10:00:00\t2026-09-01 10:00:01\t"
            "From employer\t100\t1\tEUR\tCOMPLETADO\t199\n"
            "Pago con tarjeta\tActual\t2026-09-02 10:00:00\t\t"
            "Pending shop\t-10\t0\tEUR\tPENDIENTE\t199\n",
        )
        result = parse_revolut_current(
            path,
            {"currency": "EUR", "wealthfolioAccountId": "current-account-id"},
        )
        self.assertEqual(result.checks["openingBalance"], "100")
        self.assertEqual(result.checks["skippedRows"], 1)
        self.assertEqual(
            result.checks["activityCounts"],
            {"DEPOSIT": 2, "FEE": 1},
        )
        self.assertTrue(result.warnings)

    def test_revolut_savings_reconciles_interest_reinvestment_and_withdrawal(self):
        path = self.write(
            "savings.tsv",
            "Date\tDescription\tValue, EUR\tPrice per share\tQuantity of shares\n"
            "1 ene 2025, 12:00:00\tBUY EUR Class R TEST\t2.0\t1\t2.0\n"
            "2 ene 2025, 03:00:00\tReturn PAID EUR Class R TEST\t4,707\t\t\n"
            "2 ene 2025, 03:00:00\tService Fee Charged EUR Class TEST\t-1,107\t\t\n"
            "3 ene 2025, 01:00:00\tReturn Reinvested Class R EUR TEST\t-0,36\t\t\n"
            "3 ene 2025, 12:00:00\tBUY EUR Class R TEST\t0,36\t1\t0,36\n"
            "4 ene 2025, 03:00:00\tReturn PAID EUR Class R TEST\t1,300\t\t\n"
            "4 ene 2025, 03:00:00\tService Fee Charged EUR Class TEST\t-0,300\t\t\n"
            "4 ene 2025, 12:00:00\tReturn WITHDRAWN EUR Class R TEST\t-0,10\t\t\n"
            "4 ene 2025, 12:00:00\tSELL EUR Class R TEST\t-1.000\t1\t1.000\n",
        )
        result = parse_revolut_savings(
            path,
            {"currency": "EUR", "symbol": "REV-CASH-EUR", "isin": "TEST"},
        )
        self.assertEqual(
            result.checks["activityCounts"],
            {"DEPOSIT": 1, "BUY": 2, "INTEREST": 2, "SELL": 1, "WITHDRAWAL": 1},
        )
        self.assertEqual(Decimal(result.checks["endingQuantity"]), Decimal("1000.36"))
        self.assertEqual(Decimal(result.checks["endingCash"]), Decimal("0"))
        withdrawal = next(a for a in result.activities if a.activity_type == "WITHDRAWAL")
        self.assertEqual(withdrawal.amount, Decimal("1000.10"))

    def test_xtb_reconciles_cash_and_open_positions(self):
        path = self.root / "xtb.xlsx"
        workbook = Workbook()
        closed = workbook.active
        closed.title = "Closed Positions"
        for _ in range(4):
            closed.append([])
        closed.append(["Position ID", "Open Conversion Rate", "Close Conversion Rate"])

        cash = workbook.create_sheet("Cash Operations")
        for _ in range(4):
            cash.append([])
        cash.append([
            "Type", "Instrument", "Ticker", "Category", "Time", "Amount", "ID", "Comment", "Product", "Position ID"
        ])
        cash.append(["Deposit", "", "", "", datetime(2025, 1, 1), 100, "1", "deposit", "Investment Plans", ""])
        cash.append(["Stock purchase", "Example ETF", "TEST.DE", "ETF", datetime(2025, 1, 2), -80, "2", "OPEN BUY 2/2 @ 40", "Investment Plans", "20"])
        cash.append(["Free funds interest", "", "", "", datetime(2025, 1, 3), 1, "3", "interest", "Investment Plans", ""])
        cash.append(["Subaccount transfer", "", "", "", datetime(2025, 1, 3), 5, "4", "internal", "Investment Plans", ""])
        cash.append(["Subaccount transfer", "", "", "", datetime(2025, 1, 3), -5, "5", "internal", "My Trades", ""])
        cash.append(["Total", None, None, None, None, 21, None, None, None, None])

        opened = workbook.create_sheet("Open Positions")
        for _ in range(9):
            opened.append([])
        opened.append([
            "Product", "Instrument/Position", "Ticker", "Category", "Type", "Volume", "Value", "Current price", "Open price", "Open time (UTC)"
        ])
        opened.append(["Investment Plan", "Example ETF", "TEST.DE", "ETF", "", 2, 82, "", 40, ""])
        workbook.save(path)

        result = parse_xtb(path, {"currency": "EUR"})
        self.assertEqual([a.activity_type for a in result.activities], ["DEPOSIT", "BUY", "INTEREST"])
        self.assertEqual(result.checks["endingCash"], "21")
        self.assertEqual(result.checks["ignoredSubaccountTransfers"], 2)

    def test_sabadell_reconciles_opening_balance_and_mirrors_transfer(self):
        path = self.write(
            "16092026_account.txt",
            "03/01/2026|SUPERMERCADO|03/01/2026|-10,00|140,00||SHOP-1\n"
            "02/01/2026|REMUN. CUENTA|02/01/2026|2,00|150,00||INT-1\n"
            "01/01/2026|TRASPASO A AHORROS|01/01/2026|-50,00|148,00||TR-1\n",
        )
        config = {
            "currency": "EUR",
            "wealthfolioAccountId": "main-id",
            "transferRules": [
                {
                    "pattern": "^TRASPASO A AHORROS$",
                    "targetAccount": "sabadell/ahorros",
                    "mirror": True,
                }
            ],
        }
        result = parse_sabadell(path, config)
        self.assertEqual(result.checks["statementBalance"], "140.00")
        self.assertEqual(result.checks["openingBalance"], "198.00")
        self.assertEqual(result.checks["mirroredTransfers"], 1)
        self.assertEqual(
            result.checks["activityCounts"],
            {
                "DEPOSIT": 1,
                "TRANSFER_OUT": 1,
                "TRANSFER_IN": 1,
                "INTEREST": 1,
                "WITHDRAWAL": 1,
            },
        )
        mirror = next(activity for activity in result.activities if activity.target_account)
        self.assertEqual(mirror.target_account, "sabadell/ahorros")
        opening = next(activity for activity in result.activities if activity.dedupe_key)

        rolling = self.write(
            "17092026_account.txt",
            "17/09/2026|SUPERMERCADO|17/09/2026|-5,00|95,00||SHOP-2\n",
        )
        rolling_opening = next(
            activity
            for activity in parse_sabadell(rolling, config).activities
            if activity.dedupe_key
        )
        self.assertEqual(opening.identifier, rolling_opening.identifier)

    def test_sabadell_card_maps_purchases_and_refunds(self):
        path = self.write(
            "01092026_card.txt",
            "Extracto de tarjeta\n"
            "01/09|TIENDA|MADRID|302,00 EUR\n"
            "02/09|DEVOLUCION|MADRID|-19,56 EUR\n",
        )
        result = parse_sabadell_card(path, {"currency": "EUR"})
        self.assertEqual(
            [(activity.activity_type, activity.amount) for activity in result.activities],
            [("WITHDRAWAL", Decimal("302.00")), ("CREDIT", Decimal("19.56"))],
        )

    def test_sabadell_distinguishes_same_day_equal_charges_for_mcp(self):
        path = self.write(
            "same-day.txt",
            "01/06/2026|ENVIO DE INSTANT MONEY|01/06/2026|-300,00|700,00||CARD\n"
            "01/06/2026|ENVIO DE INSTANT MONEY|01/06/2026|-300,00|400,00||CARD\n",
        )
        result = parse_sabadell(
            path,
            {"currency": "EUR", "includeOpeningBalance": False},
        )
        self.assertEqual(len(result.activities), 2)
        self.assertNotEqual(result.activities[0].comment, result.activities[1].comment)
        self.assertEqual(
            {activity.comment.rsplit("saldo ", 1)[-1] for activity in result.activities},
            {"700.00 | ref CARD", "400.00 | ref CARD"},
        )

    def test_watch_service_emits_only_new_activities(self):
        inbox = self.root / "inbox"
        account_dir = inbox / "revolut" / "stocks"
        account_dir.mkdir(parents=True)
        statement = (
            "Date\tTicker\tType\tQuantity\tPrice per share\tTotal Amount\tCurrency\tFX Rate\n"
            "2025-01-01T10:00:00Z\t\tCASH TOP-UP\t\t\tUSD 100\tUSD\t1.25\n"
        )
        config = self.root / "config.json"
        config.write_text(
            json.dumps({"accounts": {"revolut/stocks": {"parser": "revolut-stocks", "currency": "USD"}}}),
            encoding="utf-8",
        )
        environment = {
            "CONFIG_PATH": str(config),
            "INBOX_DIR": str(inbox),
            "OUTBOX_DIR": str(self.root / "outbox"),
            "PROCESSED_DIR": str(self.root / "processed"),
            "FAILED_DIR": str(self.root / "failed"),
            "STATE_DIR": str(self.root / "state"),
            "DRY_RUN": "false",
        }
        with mock.patch.dict(os.environ, environment, clear=False):
            (account_dir / "first.tsv").write_text(statement, encoding="utf-8")
            self.assertEqual(WatchService().run_once(), 0)
            (account_dir / "second.tsv").write_text(statement, encoding="utf-8")
            self.assertEqual(WatchService().run_once(), 0)
        outputs = list((self.root / "outbox").rglob("*.csv"))
        self.assertEqual(len(outputs), 1)

    def test_mcp_mapping_preserves_supported_fields_and_rejects_subtype(self):
        activity = Activity(
            date="2025-01-01T10:00:00Z",
            activity_type="BUY",
            currency="USD",
            amount=Decimal("100"),
            source_ref="source-1",
            symbol="AAPL",
            quantity=Decimal("0.5"),
            unit_price=Decimal("200"),
            fee=Decimal("1.25"),
            comment="broker reference",
        )
        row = activity_to_mcp(activity, "account-1", 7)
        self.assertEqual(row["accountId"], "account-1")
        self.assertEqual(row["lineNumber"], 7)
        self.assertEqual(row["quantity"], 0.5)
        self.assertEqual(row["fee"], 1.25)

        unsafe = Activity(
            date=activity.date,
            activity_type="CREDIT",
            currency="USD",
            amount=Decimal("1"),
            source_ref="source-2",
            subtype="REIMBURSEMENT",
        )
        with self.assertRaises(McpError):
            activity_to_mcp(unsafe, "account-1", 8)

    def test_watch_service_auto_imports_through_mcp_after_preview(self):
        inbox = self.root / "inbox"
        account_dir = inbox / "revolut" / "stocks"
        account_dir.mkdir(parents=True)
        statement = (
            "Date\tTicker\tType\tQuantity\tPrice per share\tTotal Amount\tCurrency\tFX Rate\n"
            "2025-01-01T10:00:00Z\t\tCASH TOP-UP\t\t\tUSD 100\tUSD\t1.25\n"
        )
        source = account_dir / "latest.tsv"
        source.write_text(statement, encoding="utf-8")
        config = self.root / "config.json"
        config.write_text(
            json.dumps(
                {
                    "accounts": {
                        "revolut/stocks": {
                            "parser": "revolut-stocks",
                            "currency": "USD",
                            "wealthfolioAccountId": "account-1",
                        }
                    }
                }
            ),
            encoding="utf-8",
        )
        environment = {
            "CONFIG_PATH": str(config),
            "INBOX_DIR": str(inbox),
            "OUTBOX_DIR": str(self.root / "outbox"),
            "PROCESSED_DIR": str(self.root / "processed"),
            "FAILED_DIR": str(self.root / "failed"),
            "STATE_DIR": str(self.root / "state"),
            "DRY_RUN": "false",
            "AUTO_IMPORT": "true",
        }
        client = mock.Mock()
        client.import_activities.return_value = McpImportResult(1, 0, 0, ["run-1"])
        with mock.patch.dict(os.environ, environment, clear=False), mock.patch(
            "wealthfolio_importer.service.WealthfolioMcpClient.from_environment",
            return_value=client,
        ):
            self.assertEqual(WatchService().run_once(), 0)

        client.import_activities.assert_called_once()
        _, account_id = client.import_activities.call_args.args
        self.assertEqual(account_id, "account-1")
        self.assertTrue(client.import_activities.call_args.kwargs["commit"])
        self.assertFalse(list((self.root / "outbox").rglob("*.csv")))
        self.assertFalse(source.exists())
        self.assertTrue(list((self.root / "processed").rglob("latest.tsv")))

    def test_watch_service_routes_mirrored_sabadell_transfer_to_target_account(self):
        inbox = self.root / "inbox"
        account_dir = inbox / "sabadell" / "principal"
        account_dir.mkdir(parents=True)
        source = account_dir / "01012026_account.txt"
        source.write_text(
            "01/01/2026|TRASPASO A AHORROS|01/01/2026|-50,00|150,00||TR-1\n",
            encoding="utf-8",
        )
        config = self.root / "config.json"
        config.write_text(
            json.dumps(
                {
                    "accounts": {
                        "sabadell/principal": {
                            "parser": "sabadell",
                            "currency": "EUR",
                            "wealthfolioAccountId": "main-id",
                            "transferRules": [
                                {
                                    "pattern": "^TRASPASO A AHORROS$",
                                    "targetAccount": "sabadell/ahorros",
                                    "mirror": True,
                                }
                            ],
                        },
                        "sabadell/ahorros": {
                            "parser": "sabadell",
                            "currency": "EUR",
                            "wealthfolioAccountId": "savings-id",
                        },
                    }
                }
            ),
            encoding="utf-8",
        )
        environment = {
            "CONFIG_PATH": str(config),
            "INBOX_DIR": str(inbox),
            "OUTBOX_DIR": str(self.root / "outbox"),
            "PROCESSED_DIR": str(self.root / "processed"),
            "FAILED_DIR": str(self.root / "failed"),
            "STATE_DIR": str(self.root / "state"),
            "DRY_RUN": "false",
            "AUTO_IMPORT": "true",
        }
        client = mock.Mock()
        client.import_activities.side_effect = [
            McpImportResult(2, 0, 0, ["run-main"]),
            McpImportResult(1, 0, 0, ["run-savings"]),
        ]
        with mock.patch.dict(os.environ, environment, clear=False), mock.patch(
            "wealthfolio_importer.service.WealthfolioMcpClient.from_environment",
            return_value=client,
        ):
            self.assertEqual(WatchService().run_once(), 0)

        self.assertEqual(client.import_activities.call_count, 2)
        self.assertEqual(
            [call.args[1] for call in client.import_activities.call_args_list],
            ["main-id", "savings-id"],
        )
        target_activities = client.import_activities.call_args_list[1].args[0]
        self.assertEqual([activity.activity_type for activity in target_activities], ["TRANSFER_IN"])
        self.assertFalse(source.exists())
        state = json.loads((self.root / "state" / "emitted.json").read_text())
        self.assertEqual(len(state["accounts"]["sabadell/principal"]), 3)


if __name__ == "__main__":
    unittest.main()
