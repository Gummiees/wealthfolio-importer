from __future__ import annotations

import re
import warnings
from collections import defaultdict
from datetime import datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from ..model import Activity, ConversionResult
from ..utils import iso_utc, parse_decimal, source_hash


CASH_HEADER = [
    "Type",
    "Instrument",
    "Ticker",
    "Category",
    "Time",
    "Amount",
    "ID",
    "Comment",
    "Product",
    "Position ID",
]

TRADE_PATTERN = re.compile(
    r"(?:OPEN|CLOSE)\s+(?:BUY|SELL)\s+([0-9.]+)(?:/[0-9.]+)?\s+@\s+([0-9.]+)",
    re.IGNORECASE,
)


def _decimal(value: object, field: str) -> Decimal:
    return parse_decimal(value, field)


def _load(path: Path):
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Workbook contains no default style")
        return load_workbook(path, read_only=False, data_only=True)


def _sheet_rows(sheet, header_row: int) -> tuple[list[str], list[dict[str, object]]]:
    header = [cell.value for cell in sheet[header_row]]
    while header and header[-1] is None:
        header.pop()
    rows: list[dict[str, object]] = []
    for values in sheet.iter_rows(min_row=header_row + 1, max_col=len(header), values_only=True):
        if not any(value is not None for value in values):
            continue
        rows.append(dict(zip(header, values)))
    return header, rows


def _closed_conversion_rates(workbook) -> dict[str, tuple[Decimal, Decimal]]:
    header, rows = _sheet_rows(workbook["Closed Positions"], 5)
    required = {"Position ID", "Open Conversion Rate", "Close Conversion Rate"}
    if not required.issubset(header):
        raise ValueError("La hoja Closed Positions de XTB no contiene los tipos de cambio esperados")
    result: dict[str, tuple[Decimal, Decimal]] = {}
    for row in rows:
        position_id = row.get("Position ID")
        if position_id in (None, ""):
            continue
        result[str(int(position_id)) if isinstance(position_id, float) else str(position_id)] = (
            _decimal(row["Open Conversion Rate"] or 1, "Open Conversion Rate"),
            _decimal(row["Close Conversion Rate"] or 1, "Close Conversion Rate"),
        )
    return result


def _open_position_totals(workbook) -> dict[str, Decimal]:
    sheet = workbook["Open Positions"]
    header_row = None
    for row_number in range(1, min(sheet.max_row, 30) + 1):
        values = [cell.value for cell in sheet[row_number]]
        if values[:2] == ["Product", "Instrument/Position"]:
            header_row = row_number
            break
    if header_row is None:
        raise ValueError("No se encontró la tabla de posiciones en Open Positions")
    header = [cell.value for cell in sheet[header_row]]
    while header and header[-1] is None:
        header.pop()
    result: dict[str, Decimal] = {}
    for values in sheet.iter_rows(min_row=header_row + 1, max_col=len(header), values_only=True):
        row = dict(zip(header, values))
        if row.get("Product") != "Investment Plan":
            continue
        if row.get("Type") not in (None, ""):
            continue
        ticker = str(row.get("Ticker") or "").strip()
        category = str(row.get("Category") or "").strip()
        if ticker and category:
            result[ticker] = _decimal(row["Volume"], f"Open Positions {ticker} Volume")
    return result


def parse_xtb(path: Path, config: dict) -> ConversionResult:
    account_currency = config.get("currency", "EUR")
    aliases = config.get("tickerAliases", {})
    currencies = config.get("tickerCurrencies", {})
    workbook = _load(path)
    required_sheets = {"Closed Positions", "Cash Operations", "Open Positions"}
    if set(workbook.sheetnames) != required_sheets:
        raise ValueError(
            f"Hojas XTB no reconocidas: {workbook.sheetnames}; esperadas: {sorted(required_sheets)}"
        )

    header, rows = _sheet_rows(workbook["Cash Operations"], 5)
    if header != CASH_HEADER:
        raise ValueError(f"Cabecera Cash Operations no reconocida: {header}")
    total_rows = [row for row in rows if row["Type"] == "Total"]
    if len(total_rows) != 1:
        raise ValueError("Cash Operations debe contener exactamente una fila Total")
    expected_cash = _decimal(total_rows[0]["Amount"], "Cash Operations Total")
    cash_rows = [row for row in rows if row["Type"] != "Total"]
    cash_ids = [str(row["ID"]) for row in cash_rows]
    if any(value in {"", "None"} for value in cash_ids) or len(cash_ids) != len(set(cash_ids)):
        raise ValueError("Cash Operations contiene IDs vacíos o duplicados")
    calculated_cash = sum((_decimal(row["Amount"], "Cash amount") for row in cash_rows), Decimal("0"))
    if abs(calculated_cash - expected_cash) > Decimal("0.005"):
        raise ValueError(
            f"El efectivo XTB no concilia: filas={calculated_cash}, total={expected_cash}"
        )

    rates = _closed_conversion_rates(workbook)
    activities: list[Activity] = []
    trade_quantities: dict[str, Decimal] = defaultdict(Decimal)
    cfd_rows: dict[str, list[dict[str, object]]] = defaultdict(list)
    known_types = {
        "Stock purchase",
        "Stock sell",
        "Subaccount transfer",
        "Deposit",
        "Withdrawal",
        "Free funds interest",
        "Free funds interest tax",
        "Close trade",
        "Swap",
        "Rollover",
    }

    for row in cash_rows:
        kind = str(row["Type"])
        if kind not in known_types:
            raise ValueError(f"Tipo de operación XTB desconocido: {kind!r}")
        if kind in {"Close trade", "Swap", "Rollover"}:
            position_id = str(row["Position ID"] or "")
            if not position_id:
                raise ValueError(f"{kind} sin Position ID")
            cfd_rows[position_id].append(row)
            continue
        if kind == "Subaccount transfer":
            continue

        time = row["Time"]
        if not isinstance(time, datetime):
            raise ValueError(f"Fecha XTB inválida: {time!r}")
        amount_signed = _decimal(row["Amount"], "Amount")
        cash_id = str(row["ID"])
        ref = f"xtb-cash:{cash_id}"
        common = {
            "date": iso_utc(time),
            "currency": account_currency,
            "source_ref": ref,
        }

        if kind in {"Deposit", "Withdrawal", "Free funds interest", "Free funds interest tax"}:
            activity_type = {
                "Deposit": "DEPOSIT",
                "Withdrawal": "WITHDRAWAL",
                "Free funds interest": "INTEREST",
                "Free funds interest tax": "TAX",
            }[kind]
            activities.append(
                Activity(
                    **common,
                    activity_type=activity_type,
                    amount=abs(amount_signed),
                    comment=f"XTB {kind}; cash ID {cash_id}",
                )
            )
            continue

        match = TRADE_PATTERN.fullmatch(str(row["Comment"] or "").strip())
        if not match:
            raise ValueError(f"Cash ID {cash_id}: comentario de compraventa no reconocido: {row['Comment']!r}")
        quantity = _decimal(match.group(1), "trade quantity")
        unit_price = _decimal(match.group(2), "trade price")
        source_ticker = str(row["Ticker"] or "").strip()
        if not source_ticker:
            raise ValueError(f"Cash ID {cash_id}: compraventa sin ticker")
        output_ticker = aliases.get(source_ticker, source_ticker)
        instrument_currency = currencies.get(source_ticker, account_currency)
        activity_type = "BUY" if kind == "Stock purchase" else "SELL"
        trade_quantities[source_ticker] += quantity * (1 if activity_type == "BUY" else -1)
        instrument_amount = (
            abs(amount_signed)
            if instrument_currency == account_currency
            else quantity * unit_price
        )
        fx_rate = None
        if instrument_currency != account_currency:
            position_id = str(row["Position ID"] or "")
            position_rates = rates.get(position_id)
            if position_rates:
                fx_rate = position_rates[0 if activity_type == "BUY" else 1]
            else:
                fx_rate = abs(amount_signed) / instrument_amount
        activities.append(
            Activity(
                date=common["date"],
                activity_type=activity_type,
                currency=instrument_currency,
                amount=instrument_amount,
                source_ref=ref,
                symbol=output_ticker,
                instrument_type="EQUITY",
                quantity=quantity,
                unit_price=unit_price,
                fx_rate=fx_rate,
                comment=(
                    f"XTB Stock {'purchase' if activity_type == 'BUY' else 'sell'}; "
                    f"cash ID {cash_id}; position {row['Position ID']}"
                ),
            )
        )

    for position_id, group in cfd_rows.items():
        closes = [row for row in group if row["Type"] == "Close trade"]
        if len(closes) != 1:
            raise ValueError(f"Posición CFD {position_id}: se esperaba un único Close trade")
        close = closes[0]
        net = sum((_decimal(row["Amount"], "CFD amount") for row in group), Decimal("0"))
        instrument = str(close["Instrument"] or close["Ticker"] or "CFD")
        ids = sorted(str(row["ID"]) for row in group)
        ref = source_hash("xtb-cfd", position_id, *ids)
        activities.append(
            Activity(
                date=iso_utc(close["Time"]),
                activity_type="CREDIT" if net >= 0 else "FEE",
                currency=account_currency,
                amount=abs(net),
                source_ref=f"xtb-cfd:{ref}",
                comment=(
                    f"XTB CFD realized result {instrument}; position {position_id}; "
                    "includes swap and rollover"
                ),
            )
        )

    expected_positions = _open_position_totals(workbook)
    all_tickers = set(trade_quantities) | set(expected_positions)
    position_differences: dict[str, str] = {}
    for ticker in sorted(all_tickers):
        actual = trade_quantities.get(ticker, Decimal("0"))
        expected = expected_positions.get(ticker, Decimal("0"))
        difference = actual - expected
        if abs(difference) > Decimal("0.0001"):
            position_differences[ticker] = str(difference)
    if position_differences:
        raise ValueError(f"Las posiciones XTB no concilian: {position_differences}")

    counts: dict[str, int] = defaultdict(int)
    for activity in activities:
        counts[activity.activity_type] += 1
    return ConversionResult(
        activities=activities,
        checks={
            "sourceRows": len(cash_rows),
            "ignoredSubaccountTransfers": sum(
                1 for row in cash_rows if row["Type"] == "Subaccount transfer"
            ),
            "activityCounts": dict(counts),
            "endingCash": str(expected_cash),
            "openPositions": {key: str(value) for key, value in expected_positions.items()},
        },
    ).sorted()
